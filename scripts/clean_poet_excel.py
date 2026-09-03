#!/usr/bin/env python3
"""Convert a POET therapists Excel or Google Sheets CSV into clean JSON/CSV.

Fixes auto-dates in the age column (e.g. 2026-12-03 or 3/12/2026 → 3-12),
normalizes languages/funds/settlements, maps cities to regions, and
keeps missing phone/email as empty lists without breaking card rendering.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import uuid
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent
DEFAULT_MAP = ROOT / "city_region_map.json"

HEADERS = {
    "therapist_id": "מזהה POET",
    "name": "שם המרפא/ה בעיסוק",
    "city": "עיר / יישוב",
    "age": "גילאי הילדים",
    "phone": "טלפון",
    "email": "מייל",
    "languages": "שפות טיפול",
    "kupah": 'מקבלת דרך קופ"ח? (איזו)',
    "education": 'מקבלת דרך משרד החינוך? (איזו מתי"א)',
    "online": "מקוון",
}

OPTIONAL_HEADERS = {"therapist_id", "online"}

HEADER_ALIASES = {
    "therapist id": "therapist_id",
    "poet id": "therapist_id",
    "מזהה": "therapist_id",
    "שם": "name",
    "שם המרפאה בעיסוק": "name",
    "עיר": "city",
    "יישוב": "city",
    "גילאים": "age",
    "גיל": "age",
    "טלפון": "phone",
    "מייל": "email",
    "דואל": "email",
    "שפות": "languages",
    "שפות טיפול": "languages",
    "קופה": "kupah",
    "קופת חולים": "kupah",
    'מקבלת דרך קופח? (איזו)': "kupah",
    "מתיא": "education",
    "משרד החינוך": "education",
    "טיפול מקוון": "online",
    "online": "online",
}

THERAPIST_ID_NAMESPACE = uuid.UUID("54623907-a5ee-4a6c-85bb-a460594ce5ec")

LANGUAGE_ALIASES = {
    "עברית": "hebrew",
    "hebrew": "hebrew",
    "ערבית": "arabic",
    "arabic": "arabic",
    "אנגלית": "english",
    "english": "english",
    "ספרדית": "spanish",
    "spanish": "spanish",
    "רוסית": "russian",
    "russian": "russian",
    "יידיש": "yiddish",
    "אידיש": "yiddish",
    "yiddish": "yiddish",
    "צרפתית": "french",
    "french": "french",
    "שפת סימנים": "sign_language",
    "שפת הסימנים": "sign_language",
}

LANGUAGE_LABELS = {
    "hebrew": "עברית",
    "arabic": "ערבית",
    "english": "אנגלית",
    "spanish": "ספרדית",
    "russian": "רוסית",
    "yiddish": "יידיש",
    "french": "צרפתית",
    "sign_language": "שפת סימנים",
}

FUND_LABELS = {
    "maccabi": "מכבי",
    "clalit": "כללית",
    "meuhedet": "מאוחדת",
    "leumit": "לאומית",
    "private": "קליניקה פרטית",
    "education": "משרד החינוך / מתי״א",
}

REGION_LABELS = {
    "north": "צפון",
    "sharon": "שרון",
    "center": "מרכז",
    "jerusalem": "ירושלים",
    "south": "דרום",
}

EMPTY_TOKENS = {"", "-", "—", "none", "nan", "null", "אין", "לא ידוע"}

CITY_SPLIT = re.compile(r"\s*/\s*|\s*,\s*")
PHONE_SPLIT = re.compile(r"\s*/\s*|\s*;\s*|\s+ו\s+")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"0[\d\- ]{7,14}")
AGE_RANGE_RE = re.compile(
    r"^\s*(\d+(?:\.\d+)?)\s*[-–—]\s*(\d+(?:\.\d+)?)(\s*\+)?\s*(?:\((.*?)\))?\s*$"
)
AGE_PLUS_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*\+\s*$")
ISO_DATE_RE = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?$")
SLASH_DATE_RE = re.compile(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$")


class CleanError(Exception):
    """Raised when a workbook cannot be processed."""


def load_region_map(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise CleanError(f"Region map not found: {path}")
    with path.open(encoding="utf-8") as fh:
        data = json.load(fh)
    lookup: dict[str, str] = {}
    for region, cities in data.get("regions", {}).items():
        for city in cities:
            lookup[normalize_space(city)] = region
    aliases = {normalize_space(k): normalize_space(v) for k, v in data.get("aliases", {}).items()}
    direct = {normalize_space(k): v for k, v in data.get("direct_regions", {}).items()}
    online = {normalize_space(t) for t in data.get("online_tokens", [])}
    return {"lookup": lookup, "aliases": aliases, "direct": direct, "online": online}


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\xa0", " ")).strip()


def is_empty(value: Any) -> bool:
    if value is None:
        return True
    return normalize_space(str(value)).lower() in EMPTY_TOKENS


def canon_header(value: str) -> str:
    text = normalize_space(value).replace("״", '"').replace("׳", "'")
    return text.replace('"', "").replace("'", "")


def resolve_header_key(label: str) -> str | None:
    raw = normalize_space(label)
    for key, official in HEADERS.items():
        if raw == official or canon_header(raw) == canon_header(official):
            return key
    return HEADER_ALIASES.get(raw) or HEADER_ALIASES.get(canon_header(raw))


def map_headers(labels: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, label in enumerate(labels):
        if not label:
            continue
        key = resolve_header_key(str(label))
        if key:
            found[key] = index
    missing = [HEADERS[key] for key in HEADERS if key not in found and key not in OPTIONAL_HEADERS]
    if missing:
        raise CleanError(f"Missing columns: {', '.join(missing)}")
    return found


def header_map(ws) -> dict[str, int]:
    labels = [ws.cell(1, col).value for col in range(1, (ws.max_column or 0) + 1)]
    mapped = map_headers(["" if v is None else str(v) for v in labels])
    return {key: index + 1 for key, index in mapped.items()}


def parse_age(value: Any, number_format: str | None = None) -> dict[str, Any]:
    """Return min/max/label. Excel/Sheets dates become day-month ranges."""
    result = {
        "age_min": None,
        "age_max": None,
        "age_label": None,
        "age_open_ended": False,
        "age_note": None,
        "age_raw": None if value is None else str(value),
        "age_was_date": False,
    }
    if is_empty(value):
        return result

    if isinstance(value, datetime):
        value = value.date() if not isinstance(value, date) else value.date()
    if isinstance(value, date):
        return age_from_day_month(value.day, value.month, value.isoformat())

    text = normalize_space(str(value)).lstrip("'").replace("–", "-").replace("—", "-")
    result["age_raw"] = text

    iso = ISO_DATE_RE.match(text)
    if iso:
        year, month, day = int(iso.group(1)), int(iso.group(2)), int(iso.group(3))
        if 2015 <= year <= 2035:
            return age_from_day_month(day, month, text)

    slash = SLASH_DATE_RE.match(text)
    if slash:
        first, second, year = int(slash.group(1)), int(slash.group(2)), int(slash.group(3))
        if year < 100:
            year += 2000
        if 2015 <= year <= 2035:
            day, month = first, second
            if month > 12 and 1 <= first <= 12:
                day, month = second, first
            if 1 <= month <= 12:
                return age_from_day_month(day, month, text)

    lowered = text.replace(" ", "")
    if text in {"ילדים", "ילדות", "ילדים וילדות"} or text.startswith("ילדים (") or text.startswith("ילדים("):
        result.update(age_min=0.0, age_max=18.0, age_label="ילדים (0–18)")
        return result
    if text in {"כל הגילאים", "ילדים+מבוגרים", "ילדים + מבוגרים"} or lowered in {"0-99"}:
        result.update(age_min=0.0, age_max=None, age_open_ended=True, age_label="כל הגילאים")
        return result

    plus_only = AGE_PLUS_RE.match(text)
    if plus_only:
        low = float(plus_only.group(1))
        result.update(
            age_min=low,
            age_max=None,
            age_open_ended=True,
            age_label=f"{_fmt_age(low)}+",
        )
        return result

    ranged = AGE_RANGE_RE.match(text)
    if ranged:
        low = float(ranged.group(1))
        high = float(ranged.group(2))
        open_ended = bool(ranged.group(3))
        note = ranged.group(4)
        label = f"{_fmt_age(low)}–{_fmt_age(high)}"
        if open_ended:
            label += "+"
        if note:
            result["age_note"] = note
            label += f" ({note})"
        result.update(
            age_min=low,
            age_max=None if open_ended else high,
            age_open_ended=open_ended,
            age_label=label,
        )
        return result

    result["age_label"] = text
    result["age_note"] = "unparsed"
    return result


def age_from_day_month(day: int, month: int, raw: str) -> dict[str, Any]:
    return {
        "age_min": float(day),
        "age_max": float(month),
        "age_label": f"{_fmt_age(day)}–{_fmt_age(month)}",
        "age_open_ended": False,
        "age_note": None,
        "age_raw": raw,
        "age_was_date": True,
    }


def _fmt_age(num: float | int) -> str:
    if float(num).is_integer():
        return str(int(num))
    return str(num).replace(".0", "")


def parse_online_flag(value: Any) -> bool:
    if is_empty(value):
        return False
    text = normalize_space(str(value)).lower()
    return text in {"כן", "yes", "true", "1", "מקוון", "online", "v", "x"}


def normalize_identity_text(value: Any) -> str:
    text = normalize_space("" if value is None else str(value)).lower()
    text = re.sub(r'^(ד[״"]?ר|דר)\s+', "", text)
    return re.sub(r"[\W_]+", "", text, flags=re.UNICODE)


def normalize_phone(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def normalize_therapist_id(value: Any) -> str | None:
    text = normalize_space("" if value is None else str(value)).lower()
    return text if re.fullmatch(r"poet-[0-9a-f-]{36}", text) else None


def parse_list(value: Any, splitter: re.Pattern[str]) -> list[str]:
    if is_empty(value):
        return []
    parts = [normalize_space(p) for p in splitter.split(str(value))]
    return [p for p in parts if p and p.lower() not in EMPTY_TOKENS]


def parse_phones(value: Any) -> list[str]:
    if is_empty(value):
        return []
    text = normalize_space(str(value))
    found = [normalize_space(m.group(0)) for m in PHONE_RE.finditer(text)]
    if found:
        seen: set[str] = set()
        out = []
        for phone in found:
            key = re.sub(r"\D", "", phone)
            if key not in seen:
                seen.add(key)
                out.append(phone)
        return out
    return parse_list(text, PHONE_SPLIT)


def parse_emails(value: Any) -> list[str]:
    if is_empty(value):
        return []
    found = [m.group(0).lower() for m in EMAIL_RE.finditer(str(value))]
    seen: set[str] = set()
    out = []
    for email in found:
        if email not in seen:
            seen.add(email)
            out.append(email)
    return out


def parse_languages(value: Any) -> list[str]:
    if is_empty(value):
        return ["hebrew"]
    slugs: list[str] = []
    for token in parse_list(value, re.compile(r"\s*,\s*|\s+ו\s+")):
        slug = LANGUAGE_ALIASES.get(token) or LANGUAGE_ALIASES.get(token.lower())
        if slug and slug not in slugs:
            slugs.append(slug)
    return slugs or ["hebrew"]


def parse_funds(value: Any) -> dict[str, Any]:
    text = "" if is_empty(value) else normalize_space(str(value))
    funds: list[str] = []
    reimbursement = bool(re.search(r"החזר", text))
    all_funds = bool(re.search(r"כל הקופות|כל שירותי הבריאות|כל שרותי הבריאות", text))

    checks = [
        ("maccabi", r"מכבי"),
        ("clalit", r"כללית"),
        ("meuhedet", r"מאוחדת"),
        ("leumit", r"לאומית"),
    ]
    for slug, pattern in checks:
        if re.search(pattern, text):
            funds.append(slug)

    private = bool(re.search(r"פרטי", text)) or text.startswith("לא")
    if all_funds:
        for slug in ("maccabi", "clalit", "meuhedet", "leumit"):
            if slug not in funds:
                funds.append(slug)

    if private and "private" not in funds:
        funds.append("private")
    if not funds:
        funds.append("private")

    notes = []
    if reimbursement:
        notes.append("החזרים / ביטוח משלים")
    extra = re.search(r"\((.*)\)", text)
    if extra:
        inner = extra.group(1)
        leftover = re.sub(r"מכבי|כללית|מאוחדת|לאומית|פרטי|כן|לא", " ", inner)
        leftover = normalize_space(re.sub(r"[,+/]+", " ", leftover))
        if leftover and leftover not in {"הסדר עם", "הסדרים עם"}:
            notes.append(inner)

    return {
        "funds": funds,
        "private": "private" in funds,
        "reimbursement": reimbursement,
        "all_funds": all_funds,
        "fund_raw": text or None,
        "fund_notes": " | ".join(dict.fromkeys(notes)) if notes else None,
    }


def parse_education(value: Any) -> dict[str, Any]:
    text = "" if is_empty(value) else normalize_space(str(value))
    if not text or text == "לא":
        return {"education": False, "metiv": None, "education_raw": text or None}
    metiv = None
    match = re.search(r"מתי\"?א\s*(.*)$", text)
    if match:
        metiv = normalize_space(match.group(1)).strip("() ")
        if not metiv:
            metiv = None
    return {
        "education": True,
        "metiv": metiv,
        "education_raw": text,
    }


def parse_settlements(value: Any, region_map: dict[str, Any]) -> dict[str, Any]:
    raw = None if is_empty(value) else normalize_space(str(value))
    tokens = parse_list(raw or "", CITY_SPLIT) if raw else []
    settlements: list[str] = []
    regions: list[str] = []
    unmapped: list[str] = []
    online = False

    for token in tokens:
        key = normalize_space(token)
        if key in region_map["online"] or "מקוון" in key:
            online = True
            if "מקוון" not in settlements:
                settlements.append("מקוון")
            continue
        if key in region_map["direct"]:
            region = region_map["direct"][key]
            if region not in regions:
                regions.append(region)
            if token not in settlements:
                settlements.append(token)
            continue
        canonical = region_map["aliases"].get(key, key)
        region = region_map["lookup"].get(canonical)
        display = token
        if display not in settlements:
            settlements.append(display)
        if region:
            if region not in regions:
                regions.append(region)
        else:
            unmapped.append(token)

    return {
        "settlements": settlements,
        "settlement_label": " / ".join(settlements) if settlements else (None if online else raw),
        "regions": regions,
        "online": online,
        "unmapped_cities": unmapped,
        "city_raw": raw,
    }


def build_record(row_number: int, raw: dict[str, Any], region_map: dict[str, Any]) -> dict[str, Any] | None:
    if is_empty(raw.get("name")):
        return None

    age = parse_age(raw.get("age"))
    places = parse_settlements(raw.get("city"), region_map)
    funds = parse_funds(raw.get("kupah"))
    education = parse_education(raw.get("education"))
    languages = parse_languages(raw.get("languages"))
    phones = parse_phones(raw.get("phone"))
    emails = parse_emails(raw.get("email"))
    online = places["online"] or parse_online_flag(raw.get("online"))
    if online and "מקוון" not in places["settlements"] and not places["settlements"]:
        places["settlements"] = ["מקוון"]
        places["settlement_label"] = "מקוון"

    in_person = any("מקוון" not in s for s in places["settlements"])
    if not online and not places["settlements"]:
        in_person = True

    if education["education"] and "education" not in funds["funds"]:
        funds["funds"].append("education")

    return {
        "therapist_id": normalize_therapist_id(raw.get("therapist_id")),
        "source_row": row_number,
        "name": normalize_space(str(raw["name"])),
        "settlement_label": places["settlement_label"],
        "settlements": places["settlements"],
        "regions": places["regions"],
        "region_labels": [REGION_LABELS[r] for r in places["regions"]],
        "age_min": age["age_min"],
        "age_max": age["age_max"],
        "age_label": age["age_label"],
        "age_open_ended": age["age_open_ended"],
        "phones": phones,
        "emails": emails,
        "languages": languages,
        "language_labels": [LANGUAGE_LABELS.get(s, s) for s in languages],
        "funds": funds["funds"],
        "fund_labels": [FUND_LABELS[s] for s in funds["funds"] if s in FUND_LABELS],
        "private": funds["private"],
        "reimbursement": funds["reimbursement"],
        "education": education["education"],
        "metiv": education["metiv"],
        "online": online,
        "in_person": in_person,
        "notes": " | ".join(
            p
            for p in [
                funds["fund_notes"],
                f'מתי״א {education["metiv"]}' if education["metiv"] else None,
                age["age_note"] if age["age_note"] and age["age_note"] != "unparsed" else None,
            ]
            if p
        )
        or None,
        "certified": True,
        "_age_was_date": age["age_was_date"],
        "_age_unparsed": age["age_note"] == "unparsed",
        "_age_raw": age["age_raw"],
        "_unmapped_cities": places["unmapped_cities"],
        "_fund_raw": funds["fund_raw"],
        "_education_raw": education["education_raw"],
    }


def match_known_therapist(record: dict[str, Any], known: list[dict[str, Any]]) -> tuple[str | None, str | None]:
    """Return (ID, error). Never guesses when identifying evidence conflicts."""
    name = normalize_identity_text(record.get("name"))
    emails = {str(v).strip().lower() for v in record.get("emails", []) if v}
    phones = {normalize_phone(v) for v in record.get("phones", []) if normalize_phone(v)}

    name_matches = {item["therapist_id"] for item in known if normalize_identity_text(item.get("name")) == name}
    email_matches = {
        item["therapist_id"]
        for item in known
        if emails.intersection({str(v).strip().lower() for v in item.get("emails", []) if v})
    }
    phone_matches = {
        item["therapist_id"]
        for item in known
        if phones.intersection({normalize_phone(v) for v in item.get("phones", []) if normalize_phone(v)})
    }

    if email_matches:
        compatible = email_matches
        if name_matches:
            compatible = compatible.intersection(name_matches)
            if not compatible:
                return None, "email conflicts with name"
        if len(compatible) == 1:
            return next(iter(compatible)), None
        return None, "email matches multiple therapists"

    if len(name_matches) == 1:
        candidate = next(iter(name_matches))
        if phone_matches and candidate not in phone_matches:
            return None, "phone conflicts with name"
        return candidate, None
    if len(name_matches) > 1:
        narrowed = name_matches.intersection(phone_matches)
        if len(narrowed) == 1:
            return next(iter(narrowed)), None
        return None, "name matches multiple therapists"

    if len(phone_matches) == 1:
        return next(iter(phone_matches)), None
    if len(phone_matches) > 1:
        return None, "phone matches multiple therapists"
    return None, None


def assign_therapist_ids(
    therapists: list[dict[str, Any]],
    known_items: list[dict[str, Any]],
) -> tuple[int, list[dict[str, Any]]]:
    known = [
        item for item in known_items
        if isinstance(item, dict) and normalize_therapist_id(item.get("therapist_id"))
    ]
    known_ids = {item["therapist_id"] for item in known}
    assigned_ids: set[str] = set()
    ambiguous: list[dict[str, Any]] = []
    newly_assigned = 0

    for record in therapists:
        incoming_id = normalize_therapist_id(record.get("therapist_id"))
        if incoming_id:
            if incoming_id in assigned_ids:
                ambiguous.append({
                    "row": record["source_row"],
                    "name": record["name"],
                    "reason": "duplicate therapist ID in source",
                })
                continue
            record["therapist_id"] = incoming_id
            assigned_ids.add(incoming_id)
            continue

        matched_id, error = match_known_therapist(record, known)
        if error:
            ambiguous.append({
                "row": record["source_row"],
                "name": record["name"],
                "reason": error,
            })
            continue
        if matched_id:
            if matched_id in assigned_ids:
                ambiguous.append({
                    "row": record["source_row"],
                    "name": record["name"],
                    "reason": "two incoming rows resolve to the same therapist ID",
                })
                continue
            record["therapist_id"] = matched_id
            assigned_ids.add(matched_id)
            continue

        seed = f"initial:{record['source_row']}:{normalize_identity_text(record['name'])}"
        generated_id = "poet-" + str(uuid.uuid5(THERAPIST_ID_NAMESPACE, seed))
        if generated_id in known_ids or generated_id in assigned_ids:
            ambiguous.append({
                "row": record["source_row"],
                "name": record["name"],
                "reason": "generated therapist ID collision",
            })
            continue
        record["therapist_id"] = generated_id
        assigned_ids.add(generated_id)
        newly_assigned += 1

    return newly_assigned, ambiguous


def clean_rows(
    rows: list[dict[str, Any]],
    source: str,
    sheet: str,
    region_map: dict[str, Any],
    known_items: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    therapists: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    stats: Counter[str] = Counter()

    for raw in rows:
        record = build_record(int(raw["_row"]), raw, region_map)
        if record is None:
            stats["skipped_empty_name"] += 1
            continue

        age_was_date = record.pop("_age_was_date")
        age_unparsed = record.pop("_age_unparsed")
        age_raw = record.pop("_age_raw")
        unmapped = record.pop("_unmapped_cities")
        record.pop("_fund_raw", None)
        record.pop("_education_raw", None)

        therapists.append(record)
        stats["rows"] += 1
        if age_was_date:
            stats["ages_fixed_from_date"] += 1
        if age_unparsed:
            stats["ages_unparsed"] += 1
            warnings.append({"row": record["source_row"], "type": "unparsed_age", "name": record["name"], "value": age_raw})
        if not record["phones"]:
            stats["missing_phone"] += 1
        if not record["emails"]:
            stats["missing_email"] += 1
        if not record["phones"] and not record["emails"]:
            stats["missing_both_contacts"] += 1
            warnings.append({"row": record["source_row"], "type": "no_contact", "name": record["name"]})
        if not record["settlement_label"] and not record["online"]:
            stats["missing_city"] += 1
            warnings.append({"row": record["source_row"], "type": "missing_city", "name": record["name"]})
        for city in unmapped:
            stats["unmapped_city"] += 1
            warnings.append({"row": record["source_row"], "type": "unmapped_city", "name": record["name"], "value": city})
        if record["online"]:
            stats["online"] += 1

    newly_assigned, ambiguous = assign_therapist_ids(therapists, known_items or [])
    if ambiguous:
        details = "; ".join(
            f"row {item['row']} ({item['name']}): {item['reason']}" for item in ambiguous
        )
        raise CleanError(f"Ambiguous identity match; outputs were not changed: {details}")
    stats["therapist_ids_assigned"] = newly_assigned

    return {
        "meta": {
            "source": source,
            "sheet": sheet,
            "generated_by": "scripts/clean_poet_excel.py",
            "count": len(therapists),
            "stats": dict(stats),
        },
        "warnings": warnings,
        "therapists": therapists,
    }


def read_xlsx_rows(path: Path) -> tuple[str, list[dict[str, Any]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise CleanError("Missing dependency for Excel: pip install openpyxl") from exc
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise CleanError(f"Cannot open workbook: {exc}") from exc
    ws = wb.active
    cols = header_map(ws)
    rows: list[dict[str, Any]] = []
    for row in range(2, (ws.max_row or 1) + 1):
        item = {"_row": row}
        for key in HEADERS:
            col = cols.get(key)
            item[key] = ws.cell(row, col).value if col else None
        rows.append(item)
    return ws.title, rows


def read_csv_rows(path: Path) -> tuple[str, list[dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(fh, dialect)
        try:
            labels = next(reader)
        except StopIteration as exc:
            raise CleanError("CSV file is empty") from exc
        cols = map_headers(labels)
        rows: list[dict[str, Any]] = []
        for index, values in enumerate(reader, start=2):
            item = {"_row": index}
            for key in HEADERS:
                col = cols.get(key)
                item[key] = values[col] if col is not None and col < len(values) else None
            rows.append(item)
    return path.stem, rows


def load_source(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        sheet, rows = read_csv_rows(path)
    elif suffix in {".xlsx", ".xlsm"}:
        sheet, rows = read_xlsx_rows(path)
    else:
        raise CleanError("Supported inputs: .xlsx or .csv (Google Sheets export)")
    return {"sheet": sheet, "rows": rows}


def source_kupah(item: dict[str, Any]) -> str:
    core = [label for label in item.get("fund_labels", []) if label not in {"קליניקה פרטית", "משרד החינוך / מתי״א"}]
    if core:
        return "כן (" + ", ".join(core) + ")"
    return "לא (פרטי)"


def source_education(item: dict[str, Any]) -> str:
    if not item.get("education"):
        return "לא"
    if item.get("metiv"):
        return f'כן (מתי"א {item["metiv"]})'
    return "כן"


def write_sheets_source(path: Path, therapists: list[dict[str, Any]]) -> None:
    """UTF-8 CSV ready to File→Import into Google Sheets. Age stays text via leading apostrophe."""
    fieldnames = [HEADERS[key] for key in ("therapist_id", "name", "city", "age", "phone", "email", "languages", "kupah", "education", "online")]
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for item in therapists:
            age = (item.get("age_label") or "").replace("–", "-")
            writer.writerow(
                {
                    HEADERS["therapist_id"]: item.get("therapist_id") or "",
                    HEADERS["name"]: item.get("name") or "",
                    HEADERS["city"]: item.get("settlement_label") or "",
                    HEADERS["age"]: f"'{age}" if age else "",
                    HEADERS["phone"]: " / ".join(item.get("phones") or []),
                    HEADERS["email"]: " / ".join(item.get("emails") or []),
                    HEADERS["languages"]: ", ".join(item.get("language_labels") or []),
                    HEADERS["kupah"]: source_kupah(item),
                    HEADERS["education"]: source_education(item),
                    HEADERS["online"]: "כן" if item.get("online") else "",
                }
            )


def write_csv(path: Path, therapists: list[dict[str, Any]]) -> None:
    fields = [
        "therapist_id",
        "source_row",
        "name",
        "settlement_label",
        "regions",
        "age_label",
        "age_min",
        "age_max",
        "phones",
        "emails",
        "language_labels",
        "fund_labels",
        "education",
        "metiv",
        "online",
        "in_person",
        "notes",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for item in therapists:
            row = {key: item.get(key) for key in fields}
            for key in ("regions", "phones", "emails", "language_labels", "fund_labels"):
                val = row[key]
                if isinstance(val, list):
                    row[key] = " | ".join(str(v) for v in val)
            writer.writerow(row)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Clean POET therapists Excel/CSV into JSON")
    parser.add_argument("--input", "-i", required=True, help="Path to source .xlsx or Google Sheets .csv")
    parser.add_argument("--outdir", "-o", default=str(ROOT.parent / "data"), help="Output directory")
    parser.add_argument("--map", default=str(DEFAULT_MAP), help="city_region_map.json path")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source_path = Path(args.input).expanduser().resolve()
    if not source_path.exists():
        print(f"ERROR: file not found: {source_path}", file=sys.stderr)
        return 1

    outdir = Path(args.outdir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    json_path = outdir / "therapists.json"
    known_items: list[dict[str, Any]] = []
    if json_path.exists():
        try:
            previous = json.loads(json_path.read_text(encoding="utf-8"))
            candidate_items = previous.get("therapists", previous) if isinstance(previous, dict) else previous
            if isinstance(candidate_items, list):
                known_items = candidate_items
        except (OSError, json.JSONDecodeError):
            known_items = []

    try:
        region_map = load_region_map(Path(args.map))
        loaded = load_source(source_path)
        payload = clean_rows(
            loaded["rows"],
            str(source_path),
            loaded["sheet"],
            region_map,
            known_items,
        )
    except CleanError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    csv_path = outdir / "therapists.csv"
    warn_path = outdir / "warnings.json"
    sheets_path = outdir / "poet_source_for_sheets.csv"

    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(csv_path, payload["therapists"])
    warn_path.write_text(json.dumps(payload["warnings"], ensure_ascii=False, indent=2), encoding="utf-8")
    write_sheets_source(sheets_path, payload["therapists"])

    public_data = ROOT.parent / "public" / "data"
    if (ROOT.parent / "public").exists():
        public_data.mkdir(parents=True, exist_ok=True)
        public_json = public_data / "therapists.json"
        public_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"  Hosting JSON: {public_json}")
    plugin_data = ROOT.parent / "poet-directory" / "data"
    if (ROOT.parent / "poet-directory").exists():
        plugin_data.mkdir(parents=True, exist_ok=True)
        plugin_json = plugin_data / "therapists.json"
        plugin_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"  WordPress import JSON: {plugin_json}")

    stats = payload["meta"]["stats"]
    print(f"Wrote {payload['meta']['count']} therapists")
    print(f"  JSON: {json_path}")
    print(f"  CSV:  {csv_path}")
    print(f"  Google Sheets source: {sheets_path}")
    print(f"  Warnings: {len(payload['warnings'])} → {warn_path}")
    for key in (
        "ages_fixed_from_date",
        "ages_unparsed",
        "missing_phone",
        "missing_email",
        "missing_both_contacts",
        "missing_city",
        "unmapped_city",
        "online",
    ):
        if stats.get(key):
            print(f"  {key}: {stats[key]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
